#!/usr/bin/env python3
"""Fill a copied Mother Sangha workbook without changing its form structure.

The source workbook is never written to.  Because the official files are legacy
`.xls`, LibreOffice is used only as a format bridge: xls -> temporary xlsx ->
fill cells -> xls.  The guard checks the title, header cells, and merged ranges
before and after the operation.  Unknown values are left blank on purpose.

Usage:
  fill-mother-sangha-form.py TEMPLATE.xls ROWS.json OUTPUT.xls \
      --year 2569 --level ตรี

ROWS.json must contain the array returned by mother_sangha_form_rows_v3.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


EXPECTED = {
    "ตรี": {
        "sheet": "sor5",
        "title": "บัญชีสำมะโนครัวผู้ขอเข้าสอบความรู้  ธรรมศึกษาชั้นตรี",
        "last_col": 18,
        "headers": {"A7": "เลขที่", "B7": "คำนำ", "C7": "ชื่อ", "D7": "นามสกุล", "R7": "หมายเหตุ"},
    },
    "โท": {
        "sheet": "sor6",
        "title": "บัญชีสำมะโนครัวผู้ขอเข้าสอบความรู้  ธรรมศึกษาชั้นโท",
        "last_col": 21,
        "headers": {"A7": "เลขที่", "B7": "คำนำ", "C7": "ชื่อ", "D7": "นามสกุล", "R7": "ประโยคเดิม (ธรรมศึกษาชั้นตรี)", "U7": "หมายเหตุ"},
    },
    "เอก": {
        "sheet": "sor6",
        "title": "บัญชีสำมะโนครัวผู้ขอเข้าสอบความรู้  ธรรมศึกษาชั้นเอก",
        "last_col": 21,
        "headers": {"A7": "เลขที่", "B7": "คำนำ", "C7": "ชื่อ", "D7": "นามสกุล", "R7": "ประโยคเดิม (ธรรมศึกษาชั้นโท)", "U7": "หมายเหตุ"},
    },
}

FIELD_TO_COL = {
    "form_sequence": "A", "title": "B", "first_name": "C", "last_name": "D",
    "citizen_id": "E", "dhamma_level": "F", "class_room": "G", "birth_date_be": "H",
    "organization_name": "I", "organization_subdistrict": "J", "organization_district": "K",
    "organization_province": "L", "temple_affiliation": "M", "temple_subdistrict": "N",
    "temple_district": "O", "temple_province": "P", "school_council": "Q",
    "previous_certificate_year": "R", "previous_certificate_no": "S",
    "previous_school_council": "T", "notes": "U",
}


def run_checked(command: list[str]) -> None:
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(f"command failed ({result.returncode}): {' '.join(command)}\n{result.stderr[-2000:]}")


def libreoffice() -> str:
    value = os.environ.get("SOFFICE") or shutil.which("soffice")
    if not value:
        raise RuntimeError("ไม่พบ LibreOffice/soffice สำหรับแปลงไฟล์ .xls โดยคงต้นฉบับไว้นอกการเขียน")
    return value


def convert_to_xlsx(soffice: str, source: Path, out_dir: Path) -> Path:
    run_checked([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(source)])
    result = out_dir / f"{source.stem}.xlsx"
    if not result.exists():
        raise RuntimeError(f"LibreOffice ไม่สร้างไฟล์ชั่วคราว: {result}")
    return result


def convert_to_xls(soffice: str, source_xlsx: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="dharma-xls-output-") as tmp:
        tmpdir = Path(tmp)
        run_checked([soffice, "--headless", "--convert-to", "xls:MS Excel 97", "--outdir", str(tmpdir), str(source_xlsx)])
        converted = tmpdir / f"{source_xlsx.stem}.xls"
        if not converted.exists():
            raise RuntimeError("LibreOffice ไม่สร้างไฟล์ผลลัพธ์ .xls")
        shutil.copy2(converted, output)


def check_template(ws, level: str) -> None:
    spec = EXPECTED[level]
    if ws.title != spec["sheet"]:
        raise ValueError(f"ฟอร์มระดับ{level}ต้องใช้ชีต {spec['sheet']} แต่พบ {ws.title}")
    if ws["A3"].value != spec["title"]:
        raise ValueError(f"หัวฟอร์มไม่ตรงกับระดับ{level}: {ws['A3'].value!r}")
    for address, expected in spec["headers"].items():
        if ws[address].value != expected:
            raise ValueError(f"หัวคอลัมน์ {address} ไม่ตรง: {ws[address].value!r} != {expected!r}")
    expected_merged = {str(rng) for rng in ws.merged_cells.ranges}
    if f"A3:{'R' if spec['last_col'] == 18 else 'U'}3" not in expected_merged:
        raise ValueError("merged range ของหัวฟอร์มไม่ตรงกับต้นฉบับ")


def put(cell, value) -> None:
    if value is None:
        cell.value = None
    else:
        cell.value = str(value)
        cell.number_format = "@"


def fill(template_xlsx: Path, output_xlsx: Path, rows: list[dict], year: str, level: str) -> None:
    wb = load_workbook(template_xlsx)
    ws = wb.active
    check_template(ws, level)
    spec = EXPECTED[level]
    put(ws["D4"], year)
    put(ws["F4"], "256101")
    put(ws["C5"], "โรงเรียนวัดไร่ขิงวิทยา")
    put(ws["F5"], "ไร่ขิง")
    put(ws["H5"], "สามพราน")
    put(ws["K5"], "นครปฐม")
    put(ws["N5"], "ภาคกลาง")

    for row_number, row in enumerate(rows, start=9):
        for field, column in FIELD_TO_COL.items():
            if ord(column) - ord("A") + 1 > spec["last_col"]:
                continue
            put(ws[f"{column}{row_number}"], row.get(field))
        if not row.get("form_sequence"):
            put(ws[f"A{row_number}"], row_number - 8)

    # Do not retain stale rows if a copied template already contained data.
    last_written = 8 + len(rows)
    for row_number in range(last_written + 1, ws.max_row + 1):
        for column in range(1, spec["last_col"] + 1):
            ws.cell(row_number, column).value = None
    wb.save(output_xlsx)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("template", type=Path)
    parser.add_argument("rows", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--year", required=True)
    parser.add_argument("--level", choices=EXPECTED.keys(), required=True)
    args = parser.parse_args()
    if args.template.suffix.lower() != ".xls" or args.output.suffix.lower() != ".xls":
        raise ValueError("ต้องใช้ .xls ทั้งต้นฉบับและผลลัพธ์ เพื่อรักษาฟอร์มแม่กองฯ")
    if args.template.resolve() == args.output.resolve():
        raise ValueError("ห้ามเขียนทับต้นฉบับ")
    rows = json.loads(args.rows.read_text(encoding="utf-8"))
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("ROWS.json ต้องเป็น array ของ object")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    soffice = libreoffice()
    with tempfile.TemporaryDirectory(prefix="dharma-xls-work-") as tmp:
        temp_xlsx = convert_to_xlsx(soffice, args.template, Path(tmp))
        filled_xlsx = Path(tmp) / "filled.xlsx"
        fill(temp_xlsx, filled_xlsx, rows, args.year, args.level)
        # Re-open the filled workbook and validate the exact header contract again.
        check_template(load_workbook(filled_xlsx).active, args.level)
        convert_to_xls(soffice, filled_xlsx, args.output)
    print(json.dumps({"output": str(args.output), "level": args.level, "rows": len(rows)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
